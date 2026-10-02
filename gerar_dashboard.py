#!/usr/bin/env python3
"""
Gera o dashboard HTML dos indicadores (IndicadorMET_<MES><ANO>.html) a partir
da planilha "Indicadores de Meteorologia atualizada 2026.xlsx".

POR QUE MUDOU (out/2026): o notebook antigo lia a Planilha Google ANTIGA
(SPREADSHEET_ID 1MpLz...), que parou de ser atualizada quando a planilha
virou .xlsx. Por isso setembro saiu com 100% em todas as estações.
Agora o script lê o .xlsx de verdade e CALCULA as porcentagens a partir de
D, E e F (não depende das fórmulas G/H terem sido recalculadas):

    pontualidade = (D − E) / D        consistência = (D − F) / D

Linha sem número de mensagens (D vazio ou 0) não entra (antes virava 100%).

O visual é o mesmo de antes; o modelo está em modelo_dashboard.html.
"""
import argparse
import base64
import io
import json
import os
from datetime import datetime, timezone

import openpyxl

from comum import MESES_ABREV, MESES_PT, abrir_pastas, argumentos_pastas

NOME_LOGO = 'República.jpg'
AQUI = os.path.dirname(os.path.abspath(__file__))


def _num(v):
    try:
        return float(str(v).replace(',', '.')) if v not in (None, '') else None
    except ValueError:
        return None


def ler_dados(buf_xlsx):
    from atualizar_indicadores_drive import nota_do_robo
    wb = openpyxl.load_workbook(buf_xlsx)
    dados = []
    for ws in wb.worksheets:
        cab = [str(c.value or '').strip() for c in ws[1]]
        if not cab or cab[0] != 'Estação':
            continue
        for row in ws.iter_rows(min_row=2):
            est, ano, mes = row[0].value, row[1].value, row[2].value
            d = _num(row[3].value)
            if not est or not ano or not mes or not d:
                continue
            e = _num(row[4].value) or 0
            f = _num(row[5].value) or 0
            nota_e = nota_do_robo(row[4])
            dados.append({
                'est': str(est).replace('EMS-', '').strip(),
                'ano': str(int(_num(ano))),
                'mes': str(mes).lower().strip(),
                'total': int(d),
                'pont': round(100 * (d - e) / d, 2),
                'cons': round(100 * (d - f) / d, 2),
                'prov': bool(nota_e and nota_e[0] == 'REDEMET'),
            })
    return dados


def ultimo_mes(dados):
    melhor = None
    for x in dados:
        chave = (int(x['ano']), MESES_PT.index(x['mes']) + 1 if x['mes'] in MESES_PT else 0)
        melhor = max(melhor, chave) if melhor else chave
    return melhor


def montar_html(dados, logo_bytes=None):
    src_logo = ('data:image/jpeg;base64,' + base64.b64encode(logo_bytes).decode()) if logo_bytes else ''
    modelo = open(os.path.join(AQUI, 'modelo_dashboard.html'), encoding='utf-8').read()
    agora = datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%MZ')
    return (modelo.replace('__LOGO__', src_logo)
                  .replace('__DADOS__', json.dumps(dados, ensure_ascii=False))
                  .replace('__GERADO__', agora))


def executar(pasta_ind):
    from atualizar_indicadores_drive import NOME_PLANILHA_INDICADORES
    dados = ler_dados(pasta_ind.baixar(NOME_PLANILHA_INDICADORES))
    logo = pasta_ind.baixar(NOME_LOGO).read() if pasta_ind.achar(NOME_LOGO) else None
    if logo is None:
        print(f'⚠️ Logo "{NOME_LOGO}" não encontrado na pasta; HTML sai sem o brasão.')
    ano, mes = ultimo_mes(dados)
    nome = f'IndicadorMET_{MESES_ABREV[mes - 1]}{ano}.html'
    html = montar_html(dados, logo)
    pasta_ind.enviar(nome, io.BytesIO(html.encode('utf-8')), mime='text/html')
    print(f'✅ Dashboard gerado: {nome} ({len(dados)} linhas de dados)')
    return nome


def main():
    ap = argparse.ArgumentParser()
    argumentos_pastas(ap)
    args = ap.parse_args()
    _, pasta_ind = abrir_pastas(args)
    executar(pasta_ind)


if __name__ == '__main__':
    from comum import rodar
    rodar(main)
